import pyodbc
import openpyxl
import gitignore as gi

def define_connection(database):
    connection = f'DRIVER={gi.DRIVER}; SERVER={gi.SERVER}; Trusted_Connection=yes; UID={gi.USER}; PWD={gi.PASSWORD}'
    if database:
        connection = connection + f"; Database={gi.DATABASE}"
    cnxn = pyodbc.connect(connection, autocommit=True)
    return cnxn

def connect_database(DATABASE):
    cnxn = define_connection(None)
    cursor = cnxn.cursor()
    cursor.execute("SELECT name FROM sys.databases WHERE name = ?", (DATABASE,))
    database_exist = cursor.fetchone()
    if not database_exist:
        cursor.execute(f"CREATE DATABASE {DATABASE}")
        cnxn.commit()
        print("Database doesn't exist, now creating database")
    print("Database exists")
    cnxn.close()

def check_employees_table():
    cnxn = define_connection(gi.DATABASE)
    cursor = cnxn.cursor()
    cursor.execute(f"SELECT name FROM sys.tables WHERE name = 'employees'")
    tblEmployees_exist = cursor.fetchone()
    if not tblEmployees_exist:
        print("Table employees did not exist, now creating table")
        cursor.execute(
            '''
            CREATE TABLE employees (
                id INT IDENTITY(1,1) PRIMARY KEY, 
                first_name VARCHAR(50),
                last_name VARCHAR(50),
                email VARCHAR(100) UNIQUE,
                phone_number VARCHAR(20) NULL,
                hire_date DATETIME NOT NULL,
                job_title VARCHAR(50),
                salary DECIMAL(10, 2),
                department VARCHAR(50),
                manager_id INT NULL REFERENCES employees(id) 
            )
            '''
        )
    print("Table employees exists")
    cnxn.close()

def enter_employees_table():
    cnxn = define_connection(gi.DATABASE)
    cursor = cnxn.cursor()
    wb_obj = openpyxl.load_workbook(".\Employee_Data.xlsx")
    worksheet = wb_obj["Sheet1"]
    for row in worksheet.iter_rows():
        data = []
        for cell in row:
            if(cell.value == "NULL"):
                data.append(None)
            else:
                data.append(cell.value)
        if(data[8] is not None):
            cursor.execute("SELECT id FROM employees WHERE id = ?", (data[8],))
            row = cursor.fetchone()
            if row:
                data[8] = row.id
            else:
                data[8] = None
        cursor.execute(
            '''
            insert into employees(first_name, last_name, email, phone_number, hire_date, job_title, salary, department, manager_id)
            values (?, ?, ?, ?, ?, ?, ?, ?, ?)
            '''
        , data)
        cnxn.commit()
        print(f"{data[0]} {data[1]}'s information has successfully been entered")
    print("Data has succesfully been added")
    cnxn.close()

def fetch_employees_table():
    cnxn = define_connection(gi.DATABASE)
    cursor = cnxn.cursor()
    cursor.execute("SELECT * FROM employees")
    while True:
        row = cursor.fetchone()
        if not row:
            break
        print(row)
    cnxn.close()